from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DESKTOP = Path.home() / "Desktop"
FORMULA_DOC = DESKTOP / "酒店经营模拟_核心公式需求表.docx"
HANDOFF_DOC = DESKTOP / "SWUFE_HotelSim_C方案_正式上线版交付清单_2026-04-02.docx"
BLUE_PRIMARY = RGBColor(31, 78, 121)
BLUE_SECONDARY = RGBColor(54, 96, 146)
TEXT_DARK = RGBColor(32, 32, 32)
TEXT_MUTED = RGBColor(73, 73, 73)
TABLE_HEADER_FILL = "DCE6F1"


def set_doc_defaults(doc: Document, *, margins_cm: tuple[float, float, float, float]) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.orientation = WD_ORIENTATION.PORTRAIT
    section.top_margin = Cm(margins_cm[0])
    section.bottom_margin = Cm(margins_cm[1])
    section.left_margin = Cm(margins_cm[2])
    section.right_margin = Cm(margins_cm[3])

    style_groups = [
        (["Normal"], 10.5, False),
        (["Title", "标题", "标题1", "标题 1"], 20, True),
        (["Heading 1", "标题1", "标题 1"], 15, True),
        (["Heading 2", "标题2", "标题 2"], 12.5, True),
        (["Heading 3", "标题3", "标题 3"], 11.5, True),
        (["List Bullet", "项目符号"], 10.5, False),
        (["List Number", "编号"], 10.5, False),
    ]

    for candidates, font_size, bold in style_groups:
        for style_name in candidates:
            try:
                style = doc.styles[style_name]
            except KeyError:
                continue
            style.font.name = "Microsoft YaHei"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            style.font.size = Pt(font_size)
            style.font.bold = bold


def classify_paragraph(paragraph, index: int) -> str:
    style_name = paragraph.style.name.replace(" ", "").lower()
    if index == 0:
        return "title"
    if style_name in {"heading1", "标题1"}:
        return "heading1"
    if style_name in {"heading2", "标题2"}:
        return "heading2"
    if style_name in {"heading3", "标题3"}:
        return "heading3"
    if style_name in {"listbullet", "项目符号"}:
        return "list_bullet"
    if style_name in {"listnumber", "编号"}:
        return "list_number"
    return "normal"


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_paragraph_keep_with_next(paragraph, value: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if value and keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)
    elif not value and keep_next is not None:
        p_pr.remove(keep_next)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 60, bottom: int = 60, start: int = 80, end: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for key, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)


def set_table_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def apply_table_style_if_available(table) -> None:
    for style_name in ["Table Grid", "Normal Table"]:
        try:
            table.style = style_name
            return
        except KeyError:
            continue


def set_table_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)


def format_paragraphs(doc: Document, *, title_center: bool = True) -> None:
    paragraphs = list(doc.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        paragraph_kind = classify_paragraph(paragraph, index)

        if not text:
            if paragraph_kind in {"heading1", "heading2", "heading3", "normal"}:
                remove_paragraph(paragraph)
            continue

        paragraph.paragraph_format.widow_control = True
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        paragraph.paragraph_format.line_spacing = 1.25

        if paragraph_kind == "title":
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if title_center else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(14)
            for run in paragraph.runs:
                set_run_font(run, size=20, bold=True, color=BLUE_PRIMARY)
        elif paragraph_kind == "heading1":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(8)
            set_paragraph_keep_with_next(paragraph, True)
            for run in paragraph.runs:
                set_run_font(run, size=15, bold=True, color=BLUE_PRIMARY)
        elif paragraph_kind == "heading2":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(6)
            set_paragraph_keep_with_next(paragraph, True)
            for run in paragraph.runs:
                set_run_font(run, size=12.5, bold=True, color=BLUE_SECONDARY)
        elif paragraph_kind in {"list_bullet", "list_number"}:
            paragraph.paragraph_format.left_indent = Cm(0.45)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in paragraph.runs:
                set_run_font(run, size=10.5, color=TEXT_DARK)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in paragraph.runs:
                set_run_font(run, size=10.5, color=TEXT_DARK)


def format_tables(doc: Document, *, profile: str) -> None:
    if profile == "formula":
        width_map = {
            0: [3.0, 13.2],
            1: [1.5, 3.0, 5.2, 6.5],
            2: [1.5, 3.0, 5.2, 6.5],
            3: [1.5, 3.0, 5.2, 6.5],
            4: [3.4, 12.8],
        }
    else:
        width_map = {
            0: [2.6, 2.0, 5.5, 6.5],
            1: [3.1, 2.4, 5.1, 5.8],
            2: [3.2, 2.0, 5.2, 6.0],
            3: [3.2, 2.8, 4.8, 5.6],
            4: [2.6, 1.7, 6.0, 6.1],
            5: [2.7, 2.0, 5.4, 6.3],
            6: [3.2, 3.2, 9.8],
        }

    for index, table in enumerate(doc.tables):
        apply_table_style_if_available(table)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_fixed_layout(table)
        if index in width_map:
            set_table_widths(table, width_map[index])

        repeat_table_header(table.rows[0])

        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                    paragraph.paragraph_format.line_spacing = 1.15
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if row_idx == 0
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    for run in paragraph.runs:
                        set_run_font(
                            run,
                            size=9.3 if profile == "handoff" else 9.5,
                            color=TEXT_DARK if row_idx > 0 else BLUE_PRIMARY,
                        )
                        if row_idx == 0:
                            run.font.bold = True

                if row_idx == 0:
                    shade_cell(cell, TABLE_HEADER_FILL)
                else:
                    shade_cell(cell, "FFFFFF")


def optimize_formula_doc(path: Path) -> None:
    doc = Document(path)
    set_doc_defaults(doc, margins_cm=(1.9, 1.9, 1.8, 1.8))
    format_paragraphs(doc, title_center=True)
    format_tables(doc, profile="formula")
    doc.save(path)


def optimize_handoff_doc(path: Path) -> None:
    doc = Document(path)
    set_doc_defaults(doc, margins_cm=(1.8, 1.8, 1.65, 1.65))
    format_paragraphs(doc, title_center=True)
    format_tables(doc, profile="handoff")
    doc.save(path)


def main() -> None:
    optimize_formula_doc(FORMULA_DOC)
    optimize_handoff_doc(HANDOFF_DOC)
    print(FORMULA_DOC)
    print(HANDOFF_DOC)


if __name__ == "__main__":
    main()
